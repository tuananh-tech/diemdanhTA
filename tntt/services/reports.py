"""Thống kê chuyên cần và xuất báo cáo (Excel, CSV, PDF).

CÔNG THỨC CHUYÊN CẦN (áp dụng thống nhất trong toàn hệ thống):

    Tỷ lệ chuyên cần = (Có mặt + Đi trễ) / (Có mặt + Đi trễ + Vắng có phép + Vắng không phép)

- Chỉ tính trên các phiên mà em THỰC SỰ thuộc danh sách tham dự (bản ghi được
  tạo khi mở phiên theo lớp/ngành của em tại thời điểm đó).
- "Chưa điểm danh" không nằm trong mẫu số (chưa có kết quả, không coi là vắng).
- Tùy chọn cấu hình: tính "Vắng có phép" như có mặt.
"""
from __future__ import annotations

import io
from datetime import date, timedelta

import pandas as pd
from fpdf import FPDF
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from .. import config
from .. import constants as C
from ..db import session_scope
from ..models import AttendanceRecord, AttendanceSession, Branch, ClassGroup, SchoolYear, Student
from ..security import P_REPORTS, Actor
from ..textutil import search_key
from . import org

RECORD_COLUMNS = ["session_id", "session_date", "title", "activity_type", "session_status", "student_id",
                  "code", "saint_name", "full_name", "class_id", "class_name", "branch_id", "branch", "status",
                  "method", "marked_at"]


def attendance_frame(year_id: int, actor: Actor | None = None, *, date_from: date | None = None,
                     date_to: date | None = None, branch_id: int | None = None, class_id: int | None = None,
                     activity_type: str | None = None, student_id: int | None = None) -> pd.DataFrame:
    if actor is not None:
        actor.require(P_REPORTS)
    allowed = org.allowed_class_ids(actor, year_id)
    with session_scope() as s:
        stmt = (
            select(AttendanceSession.id, AttendanceSession.session_date, AttendanceSession.title,
                   AttendanceSession.activity_type, AttendanceSession.status, Student.id, Student.code,
                   Student.saint_name, Student.full_name, AttendanceRecord.class_id, ClassGroup.name,
                   ClassGroup.branch_id, Branch.name, AttendanceRecord.status, AttendanceRecord.method,
                   AttendanceRecord.marked_at)
            .select_from(AttendanceRecord)
            .join(AttendanceSession, AttendanceSession.id == AttendanceRecord.session_id)
            .join(Student, Student.id == AttendanceRecord.student_id)
            .join(ClassGroup, ClassGroup.id == AttendanceRecord.class_id, isouter=True)
            .join(Branch, Branch.id == ClassGroup.branch_id, isouter=True)
            .where(AttendanceSession.school_year_id == year_id)
        )
        if allowed is not None:
            stmt = stmt.where(AttendanceRecord.class_id.in_(allowed or {-1}))
        if date_from:
            stmt = stmt.where(AttendanceSession.session_date >= date_from)
        if date_to:
            stmt = stmt.where(AttendanceSession.session_date <= date_to)
        if branch_id:
            stmt = stmt.where(ClassGroup.branch_id == branch_id)
        if class_id:
            stmt = stmt.where(AttendanceRecord.class_id == class_id)
        if activity_type:
            stmt = stmt.where(AttendanceSession.activity_type == activity_type)
        if student_id:
            stmt = stmt.where(AttendanceRecord.student_id == student_id)
        rows = s.execute(stmt).all()
    df = pd.DataFrame(rows, columns=RECORD_COLUMNS)
    if not df.empty:
        df["session_date"] = pd.to_datetime(df["session_date"])
        df["class_name"] = df["class_name"].fillna("(Chưa xếp lớp)")
        df["branch"] = df["branch"].fillna("(Chưa xếp ngành)")
    return df


def _attended_statuses(excused_attended: bool) -> list[str]:
    return C.ATTENDED_STATUSES + ([C.EXCUSED] if excused_attended else [])


def summarize(df: pd.DataFrame, excused_attended: bool = False) -> dict:
    counts = {st: int((df["status"] == st).sum()) if not df.empty else 0 for st in C.STATUSES}
    marked = sum(counts[s] for s in C.MARKED_STATUSES)
    attended = sum(counts[s] for s in _attended_statuses(excused_attended))
    return {
        **counts,
        "records": int(len(df)),
        "marked": marked,
        "attended": attended,
        "rate": (attended / marked * 100) if marked else None,
        "sessions": int(df["session_id"].nunique()) if not df.empty else 0,
        "students": int(df["student_id"].nunique()) if not df.empty else 0,
    }


def add_period(df: pd.DataFrame, period: str, year: SchoolYear | None = None) -> pd.DataFrame:
    df = df.copy()
    d = df["session_date"]
    if period == "day":
        df["period"] = d.dt.strftime("%Y-%m-%d")
    elif period == "week":
        df["period"] = (d - pd.to_timedelta(d.dt.day_of_week, unit="D")).dt.strftime("Tuần %Y-%m-%d")
    elif period == "month":
        df["period"] = d.dt.strftime("%Y-%m")
    elif period == "semester":
        split = pd.Timestamp(year.semester2_start) if year and year.semester2_start else None
        df["period"] = "Cả năm" if split is None else d.map(lambda x: "Học kỳ I" if x < split else "Học kỳ II")
    else:
        df["period"] = year.name if year else "Năm học"
    return df


def group_stats(df: pd.DataFrame, by: str | list[str], excused_attended: bool = False) -> pd.DataFrame:
    by = [by] if isinstance(by, str) else by
    if df.empty:
        return pd.DataFrame(columns=by + ["records", *C.STATUSES, "marked", "rate"])
    pivot = (df.assign(n=1).pivot_table(index=by, columns="status", values="n", aggfunc="sum", fill_value=0)
             .reindex(columns=C.STATUSES, fill_value=0))
    pivot["records"] = pivot[C.STATUSES].sum(axis=1)
    pivot["marked"] = pivot[C.MARKED_STATUSES].sum(axis=1)
    att = pivot[_attended_statuses(excused_attended)].sum(axis=1)
    pivot["rate"] = (att / pivot["marked"].where(pivot["marked"] > 0) * 100).round(1)
    if "session_id" in df.columns:
        pivot["sessions"] = df.groupby(by)["session_id"].nunique()
    return pivot.reset_index()


def student_stats(df: pd.DataFrame, excused_attended: bool = False) -> pd.DataFrame:
    stats = group_stats(df, ["student_id", "code", "saint_name", "full_name", "class_name", "branch"],
                        excused_attended)
    if stats.empty:
        return stats
    stats["_k"] = stats["full_name"].map(lambda n: search_key(n.split(" ")[-1]) + " " + search_key(n))
    return stats.sort_values(["class_name", "_k"]).drop(columns="_k").reset_index(drop=True)


def low_attendance(df: pd.DataFrame, threshold: float, min_marked: int = 3,
                   excused_attended: bool = False) -> pd.DataFrame:
    st = student_stats(df, excused_attended)
    if st.empty:
        return st
    return st[(st["marked"] >= min_marked) & (st["rate"] < threshold)].sort_values("rate").reset_index(drop=True)


SYMBOLS = {C.PRESENT: "✓", C.LATE: "T", C.EXCUSED: "P", C.ABSENT: "K", C.UNMARKED: "·"}


def attendance_matrix(df: pd.DataFrame) -> pd.DataFrame:
    """Sổ điểm danh: mỗi dòng 1 em, mỗi cột 1 phiên (✓ có mặt, T trễ, P phép, K không phép)."""
    if df.empty:
        return pd.DataFrame()
    sessions = (df[["session_id", "session_date", "title"]].drop_duplicates()
                .sort_values(["session_date", "session_id"]))
    labels = {r.session_id: f"{r.session_date:%d/%m} {r.title}"[:40] for r in sessions.itertuples()}
    # tránh trùng nhãn
    seen: dict[str, int] = {}
    for sid, lab in list(labels.items()):
        if lab in seen:
            seen[lab] += 1
            labels[sid] = f"{lab} ({seen[lab]})"
        else:
            seen[lab] = 1
    m = df.assign(sym=df["status"].map(SYMBOLS), col=df["session_id"].map(labels))
    out = m.pivot_table(index=["class_name", "code", "saint_name", "full_name"], columns="col", values="sym",
                        aggfunc="first", fill_value="")
    out = out.reindex(columns=[labels[s] for s in sessions["session_id"]])
    return out.reset_index()


# ---- Xuất file -----------------------------------------------------------------------
STATUS_HEADERS = {C.UNMARKED: "Chưa điểm danh", C.PRESENT: "Có mặt", C.LATE: "Đi trễ",
                  C.EXCUSED: "Vắng có phép", C.ABSENT: "Vắng không phép"}
COLUMN_LABELS = {
    "code": "Mã", "saint_name": "Tên Thánh", "full_name": "Họ và tên", "class_name": "Lớp", "branch": "Ngành",
    "records": "Số phiên thuộc đối tượng", "marked": "Đã có kết quả", "rate": "Chuyên cần (%)",
    "sessions": "Số phiên", "period": "Thời gian", "activity_type": "Loại sinh hoạt", **STATUS_HEADERS,
}


def pretty(df: pd.DataFrame, drop=("student_id", "class_id", "branch_id", "session_id")) -> pd.DataFrame:
    out = df.drop(columns=[c for c in drop if c in df.columns])
    return out.rename(columns=COLUMN_LABELS)


_FORMULA_START = ("=", "+", "@", "\t", "\r")


def _neutralize(value):
    """Chống chèn công thức khi mở file bằng Excel (CSV/Formula injection)."""
    if isinstance(value, str) and value and (
            value.startswith(_FORMULA_START) or (value.startswith("-") and not value[1:2].isdigit())):
        return "'" + value
    return value


def safe_cells(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in out.columns:
        if out[col].dtype == object or pd.api.types.is_string_dtype(out[col]):
            out[col] = out[col].map(_neutralize)
    return out


def to_csv(df: pd.DataFrame) -> bytes:
    return safe_cells(df).to_csv(index=False).encode("utf-8-sig")


def to_excel(sheets: dict[str, pd.DataFrame], *, title: str, subtitle: str, org_name: str) -> bytes:
    buf = io.BytesIO()
    navy = PatternFill("solid", fgColor="1B2A4A")
    cream = PatternFill("solid", fgColor="FBF7EE")
    thin = Side(style="thin", color="D9DCE3")
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in sheets.items():
            sheet = name[:31]
            df = safe_cells(df)
            df.to_excel(writer, index=False, sheet_name=sheet, startrow=4)
            ws = writer.sheets[sheet]
            ws["A1"] = org_name
            ws["A1"].font = Font(bold=True, size=12, color="1B2A4A")
            ws["A2"] = title
            ws["A2"].font = Font(bold=True, size=14)
            ws["A3"] = subtitle
            ws["A3"].font = Font(italic=True, color="555555")
            for cell in ws[5]:
                cell.fill = navy
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for r_idx, row in enumerate(ws.iter_rows(min_row=6, max_row=ws.max_row), start=0):
                for cell in row:
                    cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                    if r_idx % 2:
                        cell.fill = cream
            for i, col in enumerate(df.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(300).tolist()]) + 2
                ws.column_dimensions[get_column_letter(i)].width = min(max(width, 6), 40)
            ws.freeze_panes = "A6"
    return buf.getvalue()


class _ReportPDF(FPDF):
    def __init__(self, org_lines: list[str], logo_png: bytes | None, title: str, landscape: bool):
        super().__init__(orientation="L" if landscape else "P", unit="mm", format="A4")
        self.org_lines, self.logo_png, self.report_title = org_lines, logo_png, title
        fonts = config.FONTS_DIR
        self.add_font("BVP", "", str(fonts / "BeVietnamPro-Regular.ttf"))
        self.add_font("BVP", "B", str(fonts / "BeVietnamPro-Bold.ttf"))
        self.set_auto_page_break(True, margin=14)
        self.set_margins(12, 12, 12)
        self.alias_nb_pages()

    def header(self):
        x = self.l_margin
        if self.logo_png:
            self.image(io.BytesIO(self.logo_png), x, 8, 16, 16)
            x += 19
        self.set_xy(x, 8)
        self.set_font("BVP", "B", 11)
        self.set_text_color(27, 42, 74)
        self.cell(0, 5.5, self.org_lines[0], new_x="LMARGIN", new_y="NEXT")
        self.set_font("BVP", "", 8.5)
        self.set_text_color(90, 90, 90)
        for line in self.org_lines[1:]:
            self.set_x(x)
            self.cell(0, 4.2, line, new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(212, 160, 23)
        self.set_line_width(0.6)
        self.line(self.l_margin, 26, self.w - self.r_margin, 26)
        self.set_y(29)

    def footer(self):
        self.set_y(-10)
        self.set_font("BVP", "", 7.5)
        self.set_text_color(130, 130, 130)
        self.cell(0, 5, f"{self.report_title} · Trang {self.page_no()}/{{nb}}", align="C")


def to_pdf(*, title: str, subtitle: str, org_lines: list[str], logo_png: bytes | None,
           summary: dict | None, tables: list[tuple[str, pd.DataFrame]]) -> bytes:
    widest = max((len(df.columns) for _, df in tables), default=0)
    pdf = _ReportPDF(org_lines, logo_png, title, landscape=widest > 7)
    pdf.add_page()
    pdf.set_font("BVP", "B", 15)
    pdf.set_text_color(27, 42, 74)
    pdf.cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("BVP", "", 9)
    pdf.set_text_color(90, 90, 90)
    pdf.multi_cell(0, 4.5, subtitle, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    if summary:
        pdf.set_text_color(27, 42, 74)
        boxes = [("Có mặt", summary[C.PRESENT]), ("Đi trễ", summary[C.LATE]), ("Vắng có phép", summary[C.EXCUSED]),
                 ("Vắng không phép", summary[C.ABSENT]), ("Chưa điểm danh", summary[C.UNMARKED]),
                 ("Chuyên cần", f"{summary['rate']:.1f}%" if summary["rate"] is not None else "—")]
        w = (pdf.w - pdf.l_margin - pdf.r_margin) / len(boxes)
        y = pdf.get_y()
        for i, (label, value) in enumerate(boxes):
            x = pdf.l_margin + i * w
            pdf.set_fill_color(251, 247, 238)
            pdf.rect(x + 0.8, y, w - 1.6, 15, style="F", round_corners=True, corner_radius=2)
            pdf.set_xy(x, y + 1.5)
            pdf.set_font("BVP", "", 7.5)
            pdf.cell(w, 4, label, align="C")
            pdf.set_xy(x, y + 6)
            pdf.set_font("BVP", "B", 12)
            pdf.cell(w, 7, str(value), align="C")
        pdf.set_y(y + 19)
    for heading, df in tables:
        if df.empty:
            continue
        pdf.set_font("BVP", "B", 10.5)
        pdf.set_text_color(27, 42, 74)
        pdf.cell(0, 7, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("BVP", "", 7.5 if len(df.columns) > 8 else 8.5)
        pdf.set_text_color(30, 30, 30)
        data = df.copy().fillna("")
        widths = _col_widths(pdf, data)
        with pdf.table(col_widths=widths, text_align="LEFT", line_height=4.6,
                       headings_style=_heading_style(), cell_fill_color=(251, 247, 238),
                       cell_fill_mode="ROWS", borders_layout="HORIZONTAL_LINES") as table:
            table.row([str(c) for c in data.columns])
            for row in data.itertuples(index=False):
                table.row([_fmt(v) for v in row])
        pdf.ln(4)
    return bytes(pdf.output())


def _heading_style():
    from fpdf.fonts import FontFace

    return FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=(27, 42, 74))


def _fmt(v) -> str:
    if isinstance(v, float):
        return f"{v:.1f}"
    if isinstance(v, pd.Timestamp):
        return v.strftime("%d/%m/%Y")
    return str(v)


def _col_widths(pdf: FPDF, df: pd.DataFrame) -> tuple:
    raw = []
    for col in df.columns:
        sample = [str(col)] + [_fmt(v) for v in df[col].head(80).tolist()]
        raw.append(min(max(pdf.get_string_width(s) for s in sample) + 3, 70))
    return tuple(raw)


def default_range(year: SchoolYear, today: date) -> tuple[date, date]:
    end = min(today, year.end_date)
    start = max(year.start_date, end - timedelta(days=30))
    return start, end
